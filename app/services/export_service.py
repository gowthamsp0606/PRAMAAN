"""
PRAMAAN / APAAR Sentinel - Export Service
Generates sanitized CSV and Excel (.xlsx) file exports for authorized records.
Includes formula injection prevention (CWE-1236) by prefixing unsafe leading characters
('=', '+', '-', '@', '\t', '\r') with a single quote.
"""

import io
import csv
from typing import List, Any, Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExportService:
    UNSAFE_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

    @classmethod
    def sanitize_value(cls, value: Any) -> Any:
        """Prevent CSV/Excel Formula Injection (OWASP / CWE-1236)."""
        if value is None:
            return ""
        s = str(value)
        if s.startswith(cls.UNSAFE_PREFIXES):
            return f"'{s}"
        return s

    @classmethod
    def generate_csv(cls, headers: List[str], rows: List[List[Any]]) -> str:
        """Generate valid UTF-8 CSV string with sanitized cells."""
        buf = io.StringIO()
        writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([cls.sanitize_value(h) for h in headers])
        for row in rows:
            writer.writerow([cls.sanitize_value(cell) for cell in row])
        return buf.getvalue()

    @classmethod
    def generate_excel(cls, sheet_name: str, headers: List[str], rows: List[List[Any]]) -> bytes:
        """
        Generate valid .xlsx workbook using openpyxl with styled header row,
        alternating row fills, borders, and auto-adjusted column widths.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel limits sheet title to 31 chars

        # Header styling
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # Write headers
        ws.append([cls.sanitize_value(h) for h in headers])
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data rows
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        cell_font = Font(name="Calibri", size=10, color="1E293B")

        for row_idx, row in enumerate(rows, start=2):
            sanitized_row = [cls.sanitize_value(cell) for cell in row]
            ws.append(sanitized_row)
            is_alt = (row_idx % 2 == 0)
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = cell_font
                c.border = thin_border
                if is_alt:
                    c.fill = alt_fill

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        # Output to buffer
        out_buf = io.BytesIO()
        wb.save(out_buf)
        return out_buf.getvalue()
